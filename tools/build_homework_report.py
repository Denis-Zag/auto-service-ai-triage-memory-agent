from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "VPg03_Память_в_LangChain.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, *, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(0, 0, 0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def add_code_box(document: Document, text: str) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F3F4F6")
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0, 0, 0)


document = Document()
section = document.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = document.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.08

for style_name, size, before, after in (
    ("Title", 18, 0, 10),
    ("Heading 1", 14, 10, 5),
    ("Heading 2", 12, 7, 3),
):
    style = styles[style_name]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

title = document.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.add_run("Память в LangChain и других фреймворках")
subtitle = document.add_paragraph()
subtitle.add_run("Практическая работа: Auto Service AI Triage Memory Agent").bold = True

document.add_heading("Что я понял о памяти", level=1)
document.add_paragraph(
    "Память не находится внутри языковой модели. Ею управляет приложение: оно "
    "решает, что сохранить, к какому пользователю и разговору это относится и "
    "какую часть истории передать модели в следующем запросе. Поэтому хорошая "
    "память — это не просто список сообщений, а правила хранения, изоляции и "
    "очистки контекста."
)
add_bullets(
    document,
    [
        "Кратковременная память хранит состояние конкретного диалога.",
        "Долговременная память хранит выбранные факты между разными диалогами.",
        "Checkpoint позволяет продолжить процесс после нового запроса или перезапуска.",
        "client_id и thread_id нужны, чтобы не смешивать разных клиентов и разные темы.",
        "Память не отменяет проверку результата: модель может неверно понять даже сохранённый факт.",
    ],
)

document.add_heading("Сравнение трёх вариантов", level=1)
table = document.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"
headers = ["Инструмент", "Главный акцент", "Когда полезен", "Мой пример"]
for cell, value in zip(table.rows[0].cells, headers):
    set_cell_shading(cell, "D9EAF7")
    set_cell_text(cell, value, bold=True)

rows = [
    (
        "LangChain + LangGraph",
        "Диалог, tools и состояние процесса",
        "Агент задаёт вопросы и выполняет действия",
        "Помощник автосервиса",
    ),
    (
        "LlamaIndex",
        "Поиск по своим данным и документам",
        "RAG, база знаний, ответы по файлам",
        "Помощник по README и отчётам",
    ),
    (
        "Haystack",
        "Управляемый поисковый/RAG-пайплайн",
        "Нужен контроль отдельных этапов",
        "Крупная система поиска по обращениям",
    ),
]
for row_values in rows:
    cells = table.add_row().cells
    for cell, value in zip(cells, row_values):
        set_cell_text(cell, value)

document.add_paragraph(
    "Для моей практической задачи лучше всего подошли LangChain и LangGraph: "
    "нужно было не искать ответ в документах, а вести разговор, сохранять его "
    "состояние и в нужный момент вызвать Python-функцию."
)

document.add_heading("Что было до доработки", level=1)
document.add_paragraph(
    "В исходном проекте endpoint POST /triage принимал уже готовый текст обращения, "
    "классифицировал его и сохранял результат. Клиент должен был сам сразу описать "
    "всю ситуацию. Памяти и пошагового сбора сведений не было."
)

document.add_heading("Что я добавил", level=1)
add_bullets(
    document,
    [
        "endpoint POST /agent/chat для пошагового разговора;",
        "LangChain-агента со структурированным ответом Pydantic;",
        "SQLite-checkpoint LangGraph для памяти диалога;",
        "собственный tool submit_triage_request;",
        "передачу готового обращения в прежний триаж без дублирования логики;",
        "изоляцию памяти по сочетанию client_id и thread_id.",
    ],
)

document.add_heading("Как работает агент", level=1)
steps = [
    "Клиент отправляет сообщение в POST /agent/chat.",
    "Приложение получает checkpoint именно этого клиента и разговора.",
    "Агент проверяет, известны ли автомобиль, проблема и время появления.",
    "Если данных мало, задаёт один короткий вопрос и сохраняет новое состояние.",
    "Когда данных достаточно, вызывает submit_triage_request.",
    "Tool запускает существующий триаж, сохраняет заявку и возвращает ticket_id.",
    "API отвечает фиксированной структурой со статусом и результатом обработки.",
]
for index, step in enumerate(steps, 1):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.12)
    run = paragraph.add_run(f"{index}. ")
    run.bold = True
    paragraph.add_run(step)

document.add_heading("Почему разговоры не смешиваются", level=1)
document.add_paragraph(
    "Для checkpoint используется SHA-256 от сочетания client_id и thread_id. "
    "Открытые идентификаторы не добавляются в текст сообщения модели. Одинаковый "
    "thread_id у двух клиентов создаёт разные ветки памяти, а новый thread_id того "
    "же клиента начинает отдельную тему."
)

document.add_heading("Собственный tool", level=1)
document.add_paragraph(
    "Модель передаёт инструменту только сведения, которые получила в разговоре: "
    "автомобиль, симптом, время появления и дополнительные детали, если они были. "
    "Идентификатор клиента, настройки базы и API-ключ фиксирует приложение. Это "
    "снижает риск подмены служебных параметров моделью."
)

document.add_heading("Живая проверка", level=1)
document.add_paragraph(
    "Проект был собран и запущен в Docker с реальным обращением к ProxyAPI. Все "
    "примеры синтетические и не содержат персональных данных."
)

verification = document.add_table(rows=1, cols=3)
verification.style = "Table Grid"
verification.alignment = WD_TABLE_ALIGNMENT.CENTER
for cell, value in zip(verification.rows[0].cells, ["Проверка", "Действие", "Результат"]):
    set_cell_shading(cell, "D9EAF7")
    set_cell_text(cell, value, bold=True)

checks = [
    (
        "Диалог",
        "Lada Vesta, стук при повороте; затем время появления",
        "Контекст сохранён, создана заявка support",
    ),
    (
        "Перезапуск",
        "Начат разговор о Kia Rio, затем контейнер перезапущен",
        "Агент восстановил автомобиль, симптом и вопрос",
    ),
    (
        "Изоляция",
        "Тот же thread_id, но другой client_id",
        "Агент не увидел чужие данные и запросил автомобиль",
    ),
    (
        "Автотесты",
        "Полный локальный набор без сетевых обращений",
        "26 тестов пройдено",
    ),
]
for check in checks:
    cells = verification.add_row().cells
    for cell, value in zip(cells, check):
        set_cell_text(cell, value)

document.add_heading("Пример структурированного результата", level=2)
add_code_box(
    document,
    '{\n'
    '  "status": "submitted",\n'
    '  "ticket_id": 4,\n'
    '  "missing_fields": [],\n'
    '  "triage": {\n'
    '    "category": "support",\n'
    '    "confidence": "medium",\n'
    '    "escalate": false\n'
    '  }\n'
    '}',
)

document.add_heading("Что обнаружилось только в живом тесте", level=1)
document.add_paragraph(
    "После успешного вызова tool модель могла продолжить лишние циклы генерации. "
    "Это не проявлялось в обычных модульных тестах. Я сделал успешный tool-вызов "
    "конечным шагом и строю ответ API из фактически сохранённого результата. Так "
    "не создаются повторные заявки, а клиент сразу получает номер обращения."
)

document.add_heading("Ограничения учебной версии", level=1)
add_bullets(
    document,
    [
        "Модель иногда просит лишнее уточнение, даже если смысл уже понятен.",
        "Классификация и confidence не заменяют проверку человеком.",
        "Агент не ставит диагноз и не обещает цену или срок ремонта.",
        "Сейчас сохраняется состояние диалога, а не долговременный профиль клиента.",
        "Для публикации нужен постоянный серверный диск или внешнее хранилище checkpoint.",
    ],
)

document.add_heading("Вывод", level=1)
document.add_paragraph(
    "В этой работе память стала не теоретической историей сообщений, а частью "
    "реального бизнес-процесса. Агент постепенно собирает сведения, не смешивает "
    "клиентов, продолжает разговор после перезапуска и вызывает существующую "
    "функцию только тогда, когда готово обращение. Именно так память превращает "
    "одиночный запрос к модели в управляемое приложение."
)

document.save(OUTPUT)
print(OUTPUT)
